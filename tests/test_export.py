from openpyxl import load_workbook

from benchmark import export_excel
from tests.test_pipeline import make_bundle


def test_workbook_layout(tmp_path):
    bundle, _ = make_bundle(tmp_path)
    out = tmp_path / "benchmark.xlsx"
    export_excel.write(bundle, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Benchmark", "Provenance", "Exceptions", "NAV Trend"]
    bench = wb["Benchmark"]
    assert [c.value for c in bench[1]][:6] == ["Metric", "Apex Ridge", "CCLFX", "TAKIX", "GBDC", "KREF"]
    labels = [r[0].value for r in bench.iter_rows(min_row=2, max_row=10)]
    assert labels == [label for _, label in bundle.rows]  # deck rows, deck order (decision 17)
    nav_row = next(r for r in bench.iter_rows(values_only=True) if r[0] == "NAV per share")
    assert nav_row[3].startswith("$8.32")  # TAKIX
    assert nav_row[5].startswith("Book value per share: $13.04")  # KREF, exact label (decision 14)
    three_year = next(r for r in bench.iter_rows(values_only=True) if r[0].startswith("Net return, trailing 3Y"))
    assert three_year[3] == "Not reported by fund"  # TAKIX: labeled, never blank (decision 15)


def test_cells_link_to_filings_and_sheets_are_protected(tmp_path):
    bundle, _ = make_bundle(tmp_path)
    out = tmp_path / "benchmark.xlsx"
    export_excel.write(bundle, out)
    wb = load_workbook(out)
    assert all(ws.protection.sheet for ws in wb.worksheets)
    bench = wb["Benchmark"]
    competitor_cells = [c for row in bench.iter_rows(min_row=2, max_row=10, min_col=3, max_col=6) for c in row]
    # Every competitor cell, gaps included, links to the filing it came from or was looked for in (decision 21).
    assert all(c.hyperlink and c.hyperlink.target.startswith("https://www.sec.gov/") for c in competitor_cells)


def test_provenance_has_one_row_per_record(tmp_path):
    bundle, _ = make_bundle(tmp_path)
    out = tmp_path / "benchmark.xlsx"
    export_excel.write(bundle, out)
    prov = load_workbook(out)["Provenance"]
    n_records = sum(len(row) for row in bundle.table.values())
    assert prov.max_row - 1 == n_records
    header = [c.value for c in prov[1]]
    assert {"Accession", "Source quote", "Confidence", "Status"} <= set(header)
