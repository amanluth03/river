"""Excel export: the deck table plus provenance (one row per value), exceptions and the NAV trend."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .records import RunBundle, ValueRecord

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E8ECF1")
LETTER = {"high": "H", "medium": "M", "low": "L"}


def cell_text(rec: ValueRecord) -> str:
    if rec.value is None:
        return f"— ({rec.reason})" if rec.reason else "—"
    parts = [rec.display()]
    if rec.note:
        parts.append(rec.note)
    if rec.confidence:
        parts.append(f"[{LETTER[rec.confidence]}]")
    if rec.lag_quarters:
        parts.append(f"as of {rec.period_end} ({rec.lag_quarters}Q lag)")
    return " ".join(parts)


def _header(ws, names):
    ws.append(names)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEADER_FILL
    ws.freeze_panes = "B2"


def _benchmark(ws, bundle: RunBundle):
    _header(ws, ["Metric"] + bundle.funds)
    for metric, label in bundle.rows:
        row = bundle.table[metric]
        cells = [cell_text(row[f]) for f in bundle.funds]
        override = next((row[f].label for f in bundle.funds if row[f].label), None)
        ws.append([label] + cells)
        if override:
            ws.append([f"  note: {override}"] + [""] * len(bundle.funds))
    as_of_row = []
    for f in bundle.funds:
        nav = bundle.table["nav_per_share"][f]
        as_of_row.append(str(nav.period_end) if nav.period_end else "—")
    ws.append(["NAV data as of"] + as_of_row)
    ws.append([])
    for note in [
        f"As-of date {bundle.as_of}. Competitor figures from SEC EDGAR filings, each verified against the source text.",
        "[H]/[M]/[L] = confidence. See the Provenance sheet for filing, accession, section and source quote.",
        "CCLFX and TAKIX: Class I. Returns and yields on NAV basis.",
        "KREF is a mortgage REIT, not a like-for-like peer; book value shown instead of NAV. CIO to decide.",
        f"Run created {bundle.created_at:%Y-%m-%d %H:%M}, code version {bundle.code_version}.",
    ]:
        ws.append([note])
    ws.column_dimensions["A"].width = 38
    for col in "BCDEF":
        ws.column_dimensions[col].width = 30
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


PROVENANCE_COLUMNS = ["Fund", "Metric", "Value", "Unit", "Status", "Reason", "Period end", "Lag (Q)", "Form",
                      "Accession", "Filed", "Document", "URL", "Section", "Row label", "Column", "Source quote",
                      "Char offset", "SHA-256", "Method", "Confidence", "Confidence reason", "Resolution",
                      "Flags", "Note", "Other candidates"]


def _provenance_row(rec: ValueRecord) -> list:
    c = rec.citation
    others = "; ".join(f"{x.value} ({x.citation.accession}, {'verified' if x.verified else x.verify_reason})"
                       for x in rec.candidates if c is None or x.citation.accession != c.accession)
    return [rec.fund, rec.metric, str(rec.value) if rec.value is not None else "", rec.unit, rec.status, rec.reason,
            str(rec.period_end or ""), rec.lag_quarters, c and c.form, c and c.accession,
            str(c.filing_date or "") if c else "", c and c.document, c and c.url, c and c.section,
            c and c.row_label, c and c.column_label, c and c.quote, c and c.char_offset, c and c.sha256,
            rec.method, rec.confidence, rec.confidence_reason, rec.resolution, ", ".join(rec.flags), rec.note,
            others]


def write(bundle: RunBundle, path: Path) -> Path:
    wb = Workbook()
    _benchmark(wb.active, bundle)
    wb.active.title = "Benchmark"

    prov = wb.create_sheet("Provenance")
    _header(prov, PROVENANCE_COLUMNS)
    for metric, _ in bundle.rows:
        for fund in bundle.funds:
            prov.append(_provenance_row(bundle.table[metric][fund]))

    exc = wb.create_sheet("Exceptions")
    _header(exc, ["Type", "Fund", "Metric", "Message"])
    for issue in bundle.issues:
        exc.append([issue.kind, issue.fund, issue.metric or "", issue.message])
    exc.column_dimensions["D"].width = 120

    trend = wb.create_sheet("NAV Trend")
    _header(trend, ["Fund", "Quarter end", "Value", "Status / reason", "Form", "Accession", "Source quote",
                    "Confidence"])
    for fund, points in bundle.trend.items():
        for p in points:
            c = p.citation
            trend.append([fund, str(p.period_end), str(p.value) if p.value is not None else "",
                          "ok" if p.status == "ok" else f"{p.status}: {p.reason}", c and c.form,
                          c and c.accession, c and c.quote, p.confidence])
    wb.save(path)
    return Path(path)
