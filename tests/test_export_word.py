"""Word export: the same run data as the Excel, so the two can never show different numbers."""
from docx import Document

from benchmark import export_excel, export_word
from tests.test_pipeline import make_bundle


def _doc(tmp_path):
    bundle, _ = make_bundle(tmp_path)
    out = tmp_path / "benchmark.docx"
    export_word.write(bundle, out)
    return bundle, Document(out)


def test_deck_table_matches_the_excel_cells(tmp_path):
    bundle, doc = _doc(tmp_path)
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Metric"] + bundle.funds
    for r, (metric, label) in enumerate(bundle.rows, start=1):
        cells = [c.text for c in table.rows[r].cells]
        assert cells[0] == label  # deck rows, deck order
        # Every cell is exactly the Excel cell text, built by the same function from the same run.
        assert cells[1:] == [export_excel.cell_text(bundle.table[metric][f]) for f in bundle.funds]


def test_labels_gaps_and_identity(tmp_path):
    bundle, doc = _doc(tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert any(c.startswith("Book value per share: $13.04") for c in cells)
    assert "Not reported by fund" in cells and "Not in client data" in cells
    assert bundle.code_version in text and str(bundle.as_of) in text
    assert "Generated report" in text  # not a hand-edited source of record
    assert "fund accountant to confirm" in text  # Apex 1Y gap footnote


def test_every_competitor_cell_links_to_its_filing(tmp_path):
    bundle, doc = _doc(tmp_path)
    targets = {rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")}
    for metric, _ in bundle.rows:
        for fund in bundle.funds[1:]:
            cit = bundle.table[metric][fund].citation
            assert cit and cit.url in targets, (metric, fund)


def test_exceptions_and_provenance_sections(tmp_path):
    bundle, doc = _doc(tmp_path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert {"Benchmark", "Exceptions for review", "Provenance"} <= set(headings)
    exceptions = doc.tables[1]
    assert len(exceptions.rows) - 1 == len(bundle.issues)
    provenance = doc.tables[2]
    shown = [bundle.table[m][f] for m, _ in bundle.rows for f in bundle.funds[1:]]
    assert len(provenance.rows) - 1 == len(shown)
    assert any("0001725472-26-000003" in c.text for row in provenance.rows for c in row.cells)
